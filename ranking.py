"""Editable, table-based ranking slide inspired by the supplied reference."""
from io import BytesIO
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement


def filtered(frame, category='King Deal', exclude_fa=True):
    result=frame.copy()
    if category!='全部': result=result.loc[result['品類'].eq(category)].copy()
    if exclude_fa: result=result.loc[~result['通路'].eq('FA exclusive')].copy()
    return result


def ranking_slide(frames, weeks, title='', category='King Deal', exclude_fa=True):
    if not all(f.attrs.get('format')=='77Report' for f in frames):
        raise ValueError('公司排行表格版型需要 WTD77Report，請匯入原始 Excel。')
    ordered=sorted(zip(frames,weeks),key=lambda pair:pair[0].attrs['start'])
    starts=[f.attrs['start'] for f,_ in ordered]
    if len(set(starts))!=len(starts): raise ValueError('排行版型不可匯入重複週報，請每週只選一份。')
    current,week=ordered[-1]
    data=filtered(current,category,exclude_fa)
    if data.empty: raise ValueError('目前篩選沒有商品，請調整品類或 FA 條件。')
    for col in ['銷量','銷售額','單價','AUV']:
        data[col]=pd.to_numeric(data[col],errors='coerce')
        if data[col].isna().any():raise ValueError(f'{col} 有缺漏或非數字，無法產生公司排行表格。')
    if data['商品ID'].duplicated().any():raise ValueError('同週商品 ID 重複，請先確認報表。')
    top=data.sort_values(['銷量','商品ID'],ascending=[False,True]).head(10)
    total_qty=data['銷量'].sum();total_sales=data['銷售額'].sum();total_auv=data['AUV'].sum()
    delta=None
    if len(ordered)>1:
        previous=filtered(ordered[-2][0],category,exclude_fa)
        if previous.empty:raise ValueError('前週在相同篩選下沒有商品，無法比較。')
        previous_sales=pd.to_numeric(previous['銷售額'],errors='coerce')
        if previous_sales.isna().any():raise ValueError('前週銷售額有非數字，無法比較。')
        delta=total_sales-previous_sales.sum()
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    def rect(x,y,w,h,color):
        shape=slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color);shape.line.fill.background()
    def text(x,y,w,h,value,size=15,color='000000',bold=False):
        tf=slide.shapes.add_textbox(Inches(x),Inches(y), Inches(w), Inches(h)).text_frame;tf.word_wrap=True
        p=tf.paragraphs[0];p.text=value;p.font.name='Microsoft JhengHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
    rect(0,0,13.333,.87,'F4B22C');rect(.3,.14,3.2,.52,'FFF200')
    scope=f'{category}{"（排除FA）" if exclude_fa else "（含FA）"}'
    heading=title if title and title!='每週銷售洞察' else f'{scope}全品項業績追蹤'
    text(.3,.13,12.7,.6,heading,25,'000080',True)
    text(.35,.91,12.5,.35,f'{week} ｜依淨銷量排序 TOP 10 ｜All stores',13,'555555')
    table=slide.shapes.add_table(len(top)+2,7, Inches(.22),Inches(1.33), Inches(12.9), Inches(5.48)).table
    widths=[3.5,.78,1.6,2,1.55,2.05,1.42]
    for c,w in zip(table.columns,widths):c.width=Inches(w)
    def cell(row,col,value,fill='FFFFFF',size=17,color='000000',bold=False):
        c=table.cell(row,col);c.text=str(value);c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string(fill)
        c.margin_left=c.margin_right=Inches(.06);c.margin_top=c.margin_bottom=Inches(.025);c.vertical_anchor=MSO_ANCHOR.MIDDLE
        for p in c.text_frame.paragraphs:
            p.alignment=PP_ALIGN.RIGHT if col in (2,4,5,6) else PP_ALIGN.CENTER
            p.font.name='Microsoft JhengHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        tcPr=c._tc.get_or_add_tcPr()
        for edge in ('lnL','lnR','lnT','lnB'):
            line=OxmlElement('a:'+edge);line.set('w','9000');solid=OxmlElement('a:solidFill');rgb=OxmlElement('a:srgbClr');rgb.set('val','222222');solid.append(rgb);line.append(solid);tcPr.append(line)
    for j,label in enumerate(['產品','排名','單價','產品類型','銷量','業績','AUV']):cell(0,j,label,'EFB127',17,bold=True)
    for i,(_,item) in enumerate(top.iterrows(),1):
        values=[item['中文品名'],i,f"${item['單價']:,.0f}",item['產品類型'],f"{item['銷量']:,.0f}",f"${item['銷售額']:,.0f}",f"{item['AUV']:.1f}"]
        for j,value in enumerate(values):cell(i,j,value,size=14 if j in (0,3) and len(str(value))>14 else 17)
    bottom=len(top)+1
    for j in range(7):cell(bottom,j,'','D0CECE')
    cell(bottom,0,f'{scope}\n全品項合計','000000',14,'FFFFFF')
    cell(bottom,4,f'{total_qty:,.0f}','D0CECE')
    cell(bottom,5,f'${total_sales:,.0f}','D0CECE',16)
    if delta is not None:
        c=table.cell(bottom,5);p=c.text_frame.add_paragraph();p.alignment=PP_ALIGN.RIGHT
        p.text=f'（{"↑" if delta>=0 else "↓"}{abs(delta):,.0f}）';p.font.name='Microsoft JhengHei';p.font.size=Pt(13);p.font.color.rgb=RGBColor.from_string('FF0000' if delta>=0 else '008000')
    cell(bottom,6,f'{total_auv:.1f}','D0CECE')
    for i,row in enumerate(table.rows):row.height=Inches(.77 if i==bottom else 4.71/(len(top)+1))
    rect(.25,7.15,12.8,.012,'E4BD6D')
    text(.3,6.9,12.6,.28,'業績採報表 Sales；AUV 採報表值加總。合計涵蓋篩選後全品項，非僅前10名。'+('箭頭為前次匯入週業績差額。' if delta is not None else '未匯入前週，不顯示增減。'),10,'666666')
    buf=BytesIO();prs.save(buf);return buf.getvalue()
