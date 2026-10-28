import sys
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
def validate_pptx(path):
 with ZipFile(path) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith(".xml"):
    root=etree.fromstring(z.read(name))
    ids=root.xpath('//*[local-name()="cNvPr"]/@id')
    assert len(ids)==len(set(ids)),(name,ids)
    for node in root.xpath('//*[local-name()="tcPr" or local-name()="bodyPr"][@anchor]'):
     assert node.get("anchor") in {"t","ctr","b","just","dist"}

with sync_playwright() as p:
 b=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
 page=b.new_page(accept_downloads=True);errors=[];requests=[]
 page.on('pageerror',lambda e:errors.append(str(e)))
 page.on('request',lambda r:requests.append(r.url))
 page.set_content((Path(__file__).parent/'銷售簡報工作室-v2.1.html').read_text())
 page.locator('#file').set_input_files(sys.argv[1])
 page.wait_for_function("document.getElementById('status').textContent.includes('匯入成功')")
 with page.expect_download() as download:page.locator('#generate').click()
 validate_pptx(download.value.path());prs=Presentation(download.value.path());assert len(prs.slides)==1
 table=next(s.table for s in prs.slides[0].shapes if s.has_table)
 assert len(table.rows)==12
 assert table.cell(0,5).text=='業績'
 print('瀏覽器載入完整離線HTML，匯入實際Excel與原生表格PPT下載成功。')
 page.select_option('#mode','逐週銷量變化');page.locator('#generate').click()
 page.wait_for_function("document.getElementById('status').textContent.includes('至少需要兩週')")
 page.evaluate("reports.push({...reports[0],start:'2026/10/26',week:'2026/10/26–2026/11/01'})")
 with page.expect_download() as download:page.locator('#generate').click()
 validate_pptx(download.value.path());prs=Presentation(download.value.path());assert any(s.has_chart for s in prs.slides[0].shapes)
 print('雙週原生趨勢圖PPT下載成功（第二週為測試複製資料）。')
 assert not errors,errors
 assert all(u.startswith('file:') or u.startswith('blob:') for u in requests),requests
 print('無HTTP請求，無本機伺服器或Python後端。')
 b.close()
