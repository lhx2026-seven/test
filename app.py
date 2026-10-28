from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from io import BytesIO
import base64, json, webbrowser, threading, argparse
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from ranking import ranking_slide
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

ROOT = Path(__file__).parent

def read_report(name, data):
    if name.lower().endswith('.csv'):
        for encoding in ('utf-8-sig', 'cp950'):
            try:
                return pd.read_csv(BytesIO(data), encoding=encoding).fillna('')
            except UnicodeDecodeError:
                continue
        raise ValueError('CSV 編碼無法辨識，請另存 UTF-8 CSV。')
    if name.lower().endswith('.xlsx'):
        book = pd.ExcelFile(BytesIO(data))
        if 'ALL' in book.sheet_names:
            header = pd.read_excel(book, sheet_name='ALL', header=None, nrows=3)
            labels = [str(v).strip() for v in header.iloc[2]] if len(header) >= 3 else []
            if labels[:11] == ['LV1','LV2','LV3','LV4','LV5','LV6','ID','NAME','CNAME','PRICE','QTY']:
                if str(header.iloc[1, 10]).strip().lower() != 'all stores':
                    raise ValueError('77Report 找不到 All stores 區塊，請確認報表格式。')
                raw = pd.read_excel(book, sheet_name='ALL', header=2, usecols=range(16))
                valid = pd.to_numeric(raw['ID'], errors='coerce').notna()
                raw = raw.loc[valid].copy()
                ids = pd.to_numeric(raw['ID']).astype('int64').astype(str)
                names = raw['CNAME'].fillna(raw['NAME']).fillna('未命名').astype(str).str.strip()
                frame = pd.DataFrame({'品項': names + ' [' + ids + ']', '商品ID': ids,
                                      '中文品名': names, '銷量': raw['QTY'],
                                      '品類': raw['LV1'].fillna('未分類').astype(str).str.strip(),
                                      '次分類': raw['LV2'].fillna('未分類'), '銷售額': raw['Sales'],
                                      '單價': raw['PRICE'], '產品類型': raw['LV4'].fillna('未分類').replace({'Regular':'常態','Promotion':'活動','LTO':'期間限定'}),
                                      '通路': raw['LV5'].fillna('').astype(str).str.strip(), 'AUV': raw['AUV']})
                dates = [pd.to_datetime(header.iloc[1,i], errors='coerce') for i in (7,8)]
                if any(pd.isna(v) for v in dates): raise ValueError('77Report 缺少有效起訖日期。')
                frame.attrs.update(format='77Report', week=f'{dates[0]:%Y/%m/%d}–{dates[1]:%Y/%m/%d}',
                                   start=f'{dates[0]:%Y-%m-%d}', negative_rows=int((frame['銷量'] < 0).sum()))
                return frame.fillna('')
        return pd.read_excel(book).fillna('')
    raise ValueError('請匯入 .xlsx 或 .csv 檔案。')

def prepare(frame, product, quantity, category):
    if product == quantity:
        raise ValueError('品名與銷量必須選擇不同欄位。')
    values = pd.to_numeric(frame[quantity].astype(str).str.replace(',', '', regex=False), errors='coerce')
    if values.isna().any():
        raise ValueError(f'銷量欄有 {int(values.isna().sum())} 筆空白或非數字，請修正報表後重新匯入。')
    if (values < 0).any() and frame.attrs.get('format') != '77Report':
        raise ValueError('銷量包含負數；請先將退貨與銷售拆開，再匯入非負銷量。')
    names = frame[product].astype(str).str.strip()
    if names.eq('').any() or frame.empty:
        raise ValueError('報表不可為空，品名不可空白。')
    data = pd.DataFrame({'品項': names, '銷量': values})
    data['品類'] = frame[category].astype(str).str.strip().replace('', '未分類') if category else '未分類'
    if data['銷量'].sum() <= 0:
        raise ValueError('總銷量必須大於零。')
    return data

def generate(data, mode, title, period):
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor.from_string('F4F6FA')
    def box(x,y,w,h,color):
        shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color); shape.line.fill.background()
        return shape
    def text(x,y,w,h,value,size=18,color='172C45',bold=False):
        tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap=True
        for i,line in enumerate(str(value).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
            p.font.name='Microsoft JhengHei'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
    products = data.groupby('品項')['銷量'].sum().sort_values(ascending=False)
    total=float(products.sum()); top=products.index[0]
    box(0,0,13.333,.14,'20B7A8')
    text(.5,.35,11,.6,title or '每週銷售洞察',30,bold=True)
    text(.55,1,12,.4, f'{period or "匯入期間"}  /  {mode}  /  資料依匯入報表計算',12,'63748A')
    for x,label,value in [( .55,'淨銷量（匯入期間）',f'{total:,.0f}'),(4.7,'銷售品項',str(len(products))),(8.85,'第一名占比',f'{products.iloc[0]/total:.1%}')]:
        box(x,1.65,3.9,1.15,'FFFFFF');text(x+.2,1.77,3.5,.3,label,12,'63748A');text(x+.2,2.08,3.5,.55,value,28,bold=True)
    box(.55,3.05,8.05,3.8,'FFFFFF');box(8.85,3.05,3.9,3.8,'172C45')
    if mode == '逐週銷量變化':
        chart_values=data.groupby('週別', sort=False)['銷量'].sum()
        if len(chart_values)<2: raise ValueError('逐週分析至少需要兩份報表。')
        kind=XL_CHART_TYPE.LINE_MARKERS; heading='逐週總銷量變化'
        previous=float(chart_values.iloc[-2]); latest=float(chart_values.iloc[-1])
        growth=f'{(latest/previous-1):+.1%}' if previous else '前週為零，無法計算'
        note=f'最後一週銷量：{latest:,.0f}\n較前週變化：{growth}\n共比較 {len(chart_values)} 週。'
    elif mode == '品類占比':
        groups=data.groupby('品類')['銷量'].sum().sort_values(ascending=False)
        if (groups < 0).any(): raise ValueError('品類淨銷量包含負數，無法產生占比圓環圖，請改選排行或趨勢。')
        chart_values=groups.head(6).copy()
        if len(groups)>6: chart_values.loc['其他（其餘品類）']=groups.iloc[6:].sum()
        kind=XL_CHART_TYPE.DOUGHNUT; heading='品類銷量結構'
        note=f'最大品類：{groups.index[0]}\n占總銷量 {groups.iloc[0]/total:.1%}'
    else:
        low=mode == '低銷量觀察'
        chart_values=products.sort_values(ascending=low).head(8)
        kind=XL_CHART_TYPE.BAR_CLUSTERED;heading='低銷量品項 TOP 8' if low else '銷量排行 TOP 8'
        note=(f'低銷量觀察：{chart_values.index[0]}\n銷量 {chart_values.iloc[0]:,.0f}\n需搭配庫存、上市天數判斷。' if low else f'暢銷第一名：{top}\n占總銷量 {products.iloc[0]/total:.1%}\n前 3 名合計占比 {products.head(3).sum()/total:.1%}')
    text(.8,3.2,7.5,.45,heading,18,bold=True)
    cd=CategoryChartData();cd.categories=[str(v)[:35] for v in chart_values.index];cd.add_series('銷量',list(chart_values.astype(float)))
    chart=slide.shapes.add_chart(kind, Inches(.8), Inches(3.8), Inches(7.5), Inches(2.85), cd).chart
    chart.has_legend= mode == '品類占比'
    if chart.has_legend: chart.legend.position=XL_LEGEND_POSITION.RIGHT
    if kind == XL_CHART_TYPE.BAR_CLUSTERED:
        chart.category_axis.reverse_order=True
        chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=RGBColor.from_string('20B7A8')
        chart.category_axis.tick_labels.font.size=Pt(11)
    text(9.1,3.3,3.3,.5,'分析重點',21,'FFFFFF',True)
    text(9.1,4,3.25,2.45,note,17,'FFFFFF')
    text(.55,7.05,12,.25,'WEEKLY SALES  •  圖表、數據及文字皆可於 PowerPoint 編輯  •  低銷量不等於滯銷；未列入報表的品項不在分析範圍',10,'63748A')
    output=BytesIO();prs.save(output);return output.getvalue()

class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path=='/':
            self.send_response(200);self.send_header('Content-Type','text/html; charset=utf-8');self.end_headers();self.wfile.write((ROOT/'index.html').read_bytes())
        elif self.path=='/sample.csv':
            self.send_response(200);self.send_header('Content-Type','text/csv; charset=utf-8');self.end_headers();self.wfile.write((ROOT/'sample.csv').read_bytes())
        else: self.send_error(404)
    def do_POST(self):
        try:
            length=int(self.headers.get('Content-Length',0))
            if not 0<length<30_000_000: raise ValueError('檔案過大，請使用 20 MB 以下報表。')
            request=json.loads(self.rfile.read(length))
            reports=request.get('reports') or [request]
            frames=[read_report(r['name'],base64.b64decode(r['data'],validate=True)) for r in reports]
            frame=frames[0]
            common=[c for c in frame.columns if all(c in f.columns for f in frames)]
            if self.path=='/inspect':
                result=json.dumps({'columns':list(map(str,common)),'rows':sum(len(f) for f in frames),'preview':frame[common].head(5).astype(str).values.tolist(),'reports':[dict(f.attrs, name=r['name'], rows=len(f)) for f,r in zip(frames,reports)]},ensure_ascii=False).encode();ctype='application/json'
            elif self.path=='/generate' and request.get('mode')=='公司排行表格（參考第5頁）':
                weeks=[str(r.get('week') or f.attrs.get('week') or Path(r['name']).stem) for r,f in zip(reports,frames)]
                result=ranking_slide(frames,weeks,request.get('title',''),request.get('rankingCategory','King Deal'),request.get('excludeFA',True));ctype='application/vnd.openxmlformats-officedocument.presentationml.presentation'
            elif self.path=='/generate':
                parts=[]
                weeks=[str(r.get('week') or Path(r['name']).stem).strip() for r in reports]
                if len(set(weeks))!=len(weeks): raise ValueError('各份報表的週別名稱不可重複。')
                for f,week in zip(frames,weeks):
                    if request.get('excludeOthers') and f.attrs.get('format') == '77Report':
                        f=f.loc[f['品類'] != 'Others'].copy()
                    part=prepare(f,request['product'],request['quantity'],request.get('category'));part['週別']=week;parts.append(part)
                data=pd.concat(parts,ignore_index=True)
                if request['mode']=='品類占比' and not request.get('category'): raise ValueError('品類占比需要選擇品類欄位。')
                period=request['period'] or ' / '.join(weeks)
                if request.get('excludeOthers'): period += '（77Report 排除 Others）'
                result=generate(data,request['mode'],request['title'],period);ctype='application/vnd.openxmlformats-officedocument.presentationml.presentation'
            else: self.send_error(404);return
            self.send_response(200);self.send_header('Content-Type',ctype);self.end_headers();self.wfile.write(result)
        except Exception as exc:
            self.send_response(400);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(json.dumps({'error':str(exc)},ensure_ascii=False).encode())

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--no-browser',action='store_true');parser.add_argument('--port',type=int,default=8765);args=parser.parse_args()
    server=ThreadingHTTPServer(('127.0.0.1',args.port),Handler)
    if not args.no_browser: threading.Timer(.8,lambda:webbrowser.open(f'http://127.0.0.1:{args.port}')).start()
    print(f'銷售簡報工作室已啟動，連接埠 {args.port}；Ctrl+C 結束。',flush=True);server.serve_forever()
