import threading, json, base64, urllib.request
from io import BytesIO
from http.server import ThreadingHTTPServer
from pptx import Presentation
from app import Handler, ROOT
server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
threading.Thread(target=server.serve_forever,daemon=True).start()
base=f'http://127.0.0.1:{server.server_port}'
try:
    assert '銷售簡報工作室' in urllib.request.urlopen(base).read().decode()
    encoded=base64.b64encode((ROOT/'sample.csv').read_bytes()).decode()
    body={'reports':[{'name':'01.csv','data':encoded,'week':'第一週'},{'name':'02.csv','data':encoded,'week':'第二週'}]}
    def post(path,payload):
        return urllib.request.urlopen(urllib.request.Request(base+path,json.dumps(payload).encode(),{'Content-Type':'application/json'}))
    assert json.load(post('/inspect',body))['rows']==20
    body.update(product='品名',quantity='銷量',category='品類',mode='逐週銷量變化',title='跨週測試',period='兩週')
    result=post('/generate',body).read()
    prs=Presentation(BytesIO(result))
    chart=next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values)==[2205,2205]
    assert [c.label for c in chart.plots[0].categories]==['第一週','第二週']
    (ROOT/'範例逐週分析.pptx').write_bytes(result)
    print('HTTP 首頁、多檔匯入、週趨勢 PPT 與內嵌圖表數據驗證通過。')
finally:server.shutdown();server.server_close()
