import unittest
from io import BytesIO
from pptx import Presentation
import pandas as pd
from ranking import ranking_slide, filtered
from app import read_report, prepare, generate, ROOT

class WorkflowTests(unittest.TestCase):
    def test_csv_and_all_slides(self):
        frame=read_report('sample.csv',(ROOT/'sample.csv').read_bytes())
        data=prepare(frame,'品名','銷量','品類')
        self.assertEqual(data['銷量'].sum(),2205)
        for mode in ('銷量排行','品類占比','低銷量觀察'):
            output=generate(data,mode,'測試週報','2026/10/19–25')
            prs=Presentation(BytesIO(output));self.assertEqual(len(prs.slides),1)
            charts=[s.chart for s in prs.slides[0].shapes if s.has_chart]
            self.assertEqual(len(charts),1)
            self.assertTrue(any(s.has_text_frame and '2,205' in s.text for s in prs.slides[0].shapes))
            self.assertGreater(len(charts[0].series[0].values),0)
    def test_excel_and_aggregation(self):
        frame=pd.DataFrame({'品名':['A','A','B'],'銷量':['1,000','20','5']})
        buf=BytesIO();frame.to_excel(buf,index=False)
        data=prepare(read_report('test.xlsx',buf.getvalue()),'品名','銷量',None)
        self.assertEqual(data.groupby('品項')['銷量'].sum()['A'],1020)
    def test_wtd_format(self):
        rows=[['metadata']+[None]*15,
              ['BK_77Report',None,None,None,None,None,'Date',pd.Timestamp('2026-10-19'),pd.Timestamp('2026-10-25'),None,'All stores']+[None]*5,
              ['LV1','LV2','LV3','LV4','LV5','LV6','ID','NAME','CNAME','PRICE','QTY','Sales','Sales%','Cost','AUV','Incidence'],
              ['CORE','Beef',None,None,None,None,1,'English','同名',100,10,1000,0,0,0,0],
              ['CORE','Beef',None,None,None,None,2,'English','同名',100,-1,-100,0,0,0,0],
              [None]*10+[9999,0,0,0,0,0]]
        buf=BytesIO()
        with pd.ExcelWriter(buf,engine='openpyxl') as writer:
            pd.DataFrame(rows).to_excel(writer,sheet_name='ALL',header=False,index=False)
        frame=read_report('week.xlsx',buf.getvalue())
        self.assertEqual(len(frame),2)
        self.assertEqual(frame.attrs['week'],'2026/10/19–2026/10/25')
        self.assertEqual(frame.attrs['negative_rows'],1)
        data=prepare(frame,'品項','銷量','品類')
        self.assertEqual(data['銷量'].sum(),9)
        self.assertEqual(data['品項'].nunique(),2)
        output=ranking_slide([frame],['測試週'],category='CORE',exclude_fa=True)
        table=next(shape.table for shape in Presentation(BytesIO(output)).slides[0].shapes if shape.has_table)
        self.assertEqual(table.cell(0,5).text,'業績')
        self.assertEqual(table.cell(3,4).text,'9')
        self.assertEqual(table.cell(3,5).text,'$900')
        self.assertEqual(table.cell(1,0).text,'同名')
        second=frame.copy();second.attrs['start']='2026-10-26'
        output=ranking_slide([frame,second],['第一週','第二週'],category='CORE')
        table=next(shape.table for shape in Presentation(BytesIO(output)).slides[0].shapes if shape.has_table)
        self.assertIn('↑0',table.cell(3,5).text)
        with self.assertRaises(ValueError):ranking_slide([frame,frame],['第一週','重複週'],category='CORE')
        frame.loc[frame.index[0],'通路']='FA exclusive'
        self.assertEqual(len(filtered(frame,'CORE',True)),1)

    def test_mapping_changes_analysis(self):
        f=pd.DataFrame({'商品':['A','B'],'別名':['X','Y'],'件數':[2,8],'份數':[30,70],'分類':['甲','乙']})
        first=prepare(f,'商品','件數','分類')
        second=prepare(f,'別名','份數',None)
        self.assertEqual(first['銷量'].sum(),10)
        self.assertEqual(second['銷量'].sum(),100)
        self.assertEqual(second['品項'].tolist(),['X','Y'])
        self.assertEqual(first['品類'].tolist(),['甲','乙'])
        self.assertEqual(second['品類'].tolist(),['未分類','未分類'])

    def test_invalid_quantities(self):
        for values in [['x'],[-1],[0],['']]:
            with self.assertRaises(ValueError): prepare(pd.DataFrame({'p':['A'],'q':values}),'p','q',None)
    def test_empty_product(self):
        with self.assertRaises(ValueError):prepare(pd.DataFrame({'p':[' '],'q':[2]}),'p','q',None)

if __name__=='__main__':unittest.main()
